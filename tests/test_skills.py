"""
Skills: SKILL.md parsing, discovery and matching, importing, structured runs,
and the bundled document skills.
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import jsonschema
import pytest

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT_DIR)

from agents.cua_cli.agent import CLIAgent  # noqa: E402
from agents.skills.importer import SkillImportError, import_skills, install_skill_dir, resolve_source  # noqa: E402
from agents.skills.manifest import SkillParseError, parse_skill_file  # noqa: E402
from agents.skills.registry import BUNDLED_SKILLS_DIR, SkillRegistry  # noqa: E402
from agents.skills.structured import (  # noqa: E402
    extract_json_object,
    generate_valid_spec,
    resolve_output_path,
    run_structured_skill,
)

DOCUMENT_SKILLS = ("pptx", "docx", "pdf", "xlsx")


def _example_spec(skill: str) -> dict:
    body = (BUNDLED_SKILLS_DIR / skill / "SKILL.md").read_text(encoding="utf-8")
    return json.loads(body.split("```json", 1)[1].split("```", 1)[0])


def _write_skill(root: Path, name: str, description: str, *, extra_frontmatter: str = "", body: str = "Do the thing.") -> Path:
    skill_dir = root / name
    skill_dir.mkdir(parents=True)
    (skill_dir / "SKILL.md").write_text(
        f"---\nname: {name}\ndescription: {description}\n{extra_frontmatter}---\n\n{body}\n", encoding="utf-8"
    )
    return skill_dir


@pytest.fixture
def user_dir(tmp_path, monkeypatch) -> Path:
    path = tmp_path / "user-skills"
    path.mkdir()
    monkeypatch.setenv("JARVIS_SKILLS_DIR", str(path))
    return path


# Bundled document skills -----------------------------------------------------


@pytest.mark.parametrize("skill", DOCUMENT_SKILLS)
def test_bundled_skill_examples_match_their_schemas(skill: str) -> None:
    schema = json.loads((BUNDLED_SKILLS_DIR / skill / "schema.json").read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator.check_schema(schema)
    jsonschema.validate(_example_spec(skill), schema)


@pytest.mark.parametrize("skill", DOCUMENT_SKILLS)
async def test_bundled_skill_builds_its_example(skill: str, tmp_path, monkeypatch, user_dir) -> None:
    monkeypatch.setenv("JARVIS_OUTPUT_DIR", str(tmp_path / "out"))
    entry = SkillRegistry.discover().get(skill)
    assert entry is not None and entry.tier == "bundled" and entry.manifest.is_structured

    async def _generate(system, user, force):
        assert "JSON Schema" in system and "make it" in user
        return json.dumps(_example_spec(skill))

    result = await run_structured_skill(entry, "make it", generate=_generate)

    assert result["success"] is True, result
    path = Path(result["output_file_path"])
    assert path.suffix == f".{skill}" and path.parent == tmp_path / "out"
    assert result["tool_calls"][0] == {
        "tool_name": "write_file",
        "tool_id": f"skill-{skill}",
        "parameters": {"file_path": str(path)},
        "result": result["tool_calls"][0]["result"],
        "status": "success",
    }
    if skill == "pptx":
        from pptx import Presentation

        assert len(Presentation(str(path)).slides) == 5
    elif skill == "docx":
        from docx import Document

        assert Document(str(path)).paragraphs[0].text == "Weekly Lab Report"
    elif skill == "pdf":
        assert path.read_bytes().startswith(b"%PDF")
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(str(path))["October Budget"]
        assert sheet["D2"].value == "=B2-C2" and sheet["B5"].value == "=SUM(B2:B4)"


# Parsing --------------------------------------------------------------------


def test_parse_rejects_invalid_names_and_missing_description(tmp_path) -> None:
    bad_name = _write_skill(tmp_path, "Bad_Name", "x")
    with pytest.raises(SkillParseError):
        parse_skill_file(bad_name / "SKILL.md")
    no_description = tmp_path / "nodesc"
    no_description.mkdir()
    (no_description / "SKILL.md").write_text("---\nname: nodesc\n---\nbody\n", encoding="utf-8")
    with pytest.raises(SkillParseError):
        parse_skill_file(no_description / "SKILL.md")


def test_parse_reads_openjarvis_capabilities(tmp_path) -> None:
    skill = _write_skill(
        tmp_path, "runner", "Runs things",
        extra_frontmatter="metadata:\n  openjarvis:\n    required_capabilities: [shell:execute, network:fetch]\n",
    )
    manifest = parse_skill_file(skill / "SKILL.md")
    assert manifest.required_capabilities == ("shell:execute", "network:fetch")
    assert manifest.dangerous_capabilities == ("shell:execute",)
    assert manifest.body == "Do the thing."


# Discovery and matching ---------------------------------------------------------


@pytest.mark.parametrize(
    ("task", "expected"),
    [
        ("make a 6 slide presentation on solar energy", "pptx"),
        ("create a PowerPoint about our trip", "pptx"),
        ("write a word document about my internship", "docx"),
        ("prepare a PDF handout on first aid", "pdf"),
        ("make an excel budget for october", "xlsx"),
        ("create a spreadsheet tracking my expenses and save it as a pdf", "xlsx"),
        ("use the pdf skill for my notes", "pdf"),
        ("/docx meeting notes", "docx"),
        ("open the presentation on my desktop", None),
        ("what is a pdf", None),
    ],
)
def test_structured_skill_matching(task: str, expected: str | None, user_dir) -> None:
    match = SkillRegistry.discover().structured_skill_for(task)
    assert (match.name if match else None) == expected


def test_tiers_and_precedence(user_dir) -> None:
    _write_skill(user_dir, "pptx", "My own deck style")
    imported = _write_skill(user_dir / "imported" / "hermes", "arxiv", "Search arXiv papers by keyword")
    (imported / ".jarvis-source.json").write_text(json.dumps({"source": "hermes", "scripts_approved": False}), encoding="utf-8")
    shadowed = _write_skill(user_dir / "imported" / "other", "docx", "Imported docx")
    (shadowed / ".jarvis-source.json").write_text("{}", encoding="utf-8")

    registry = SkillRegistry.discover()

    assert registry.get("pptx").tier == "workspace"
    assert registry.get("docx").tier == "bundled"
    arxiv = registry.get("arxiv")
    assert (arxiv.tier, arxiv.source, arxiv.can_run_scripts) == ("imported", "hermes", False)


def test_relevant_instruction_skills(user_dir) -> None:
    _write_skill(user_dir, "arxiv", "Search arXiv papers by keyword, author, or category")
    _write_skill(user_dir, "git-hygiene", "Clean up git branches and stale remotes")
    registry = SkillRegistry.discover()

    assert [entry.name for entry in registry.relevant_skills("find recent arxiv papers about diffusion")] == ["arxiv"]
    assert registry.relevant_skills("open notepad") == []
    assert registry.explicit_skill("use the git-hygiene skill on this repo").name == "git-hygiene"


# Importing --------------------------------------------------------------------


def test_import_skips_scripts_unless_approved(tmp_path, user_dir) -> None:
    source = _write_skill(tmp_path / "src", "arxiv", "Search arXiv papers")
    (source / "scripts").mkdir()
    (source / "scripts" / "search.py").write_text("print('hi')\n", encoding="utf-8")
    (source / "references").mkdir()
    (source / "references" / "api.md").write_text("docs\n", encoding="utf-8")

    [outcome] = import_skills(str(source.parent))
    target = user_dir / "imported" / "local" / "arxiv"
    assert outcome.imported and outcome.target == target
    assert (target / "references" / "api.md").exists() and not (target / "scripts").exists()
    assert json.loads((target / ".jarvis-source.json").read_text())["scripts_approved"] is False

    [again] = import_skills(str(source.parent))
    assert not again.imported and "already installed" in again.skipped_reason

    [forced] = import_skills(str(source.parent), with_scripts=True, force=True)
    assert forced.scripts_imported and (target / "scripts" / "search.py").exists()
    assert SkillRegistry.discover().get("arxiv").can_run_scripts


def test_import_refuses_native_symlinks(tmp_path, user_dir) -> None:
    linked = _write_skill(tmp_path / "a", "linked", "Has a symlink")
    try:
        (linked / "secret").symlink_to(tmp_path, target_is_directory=True)
    except OSError as exc:
        if getattr(exc, "winerror", None) == 1314:
            pytest.skip("Windows account lacks the privilege to create native symlinks")
        raise
    outcome = install_skill_dir(linked, source="local", origin="x")
    assert not outcome.imported and "symlink" in outcome.skipped_reason


def test_import_rejects_symlink_before_copying(tmp_path, user_dir, monkeypatch) -> None:
    linked = _write_skill(tmp_path / "a", "linked", "Has a symlink")
    link = linked / "secret"
    link.touch()
    original_is_symlink = Path.is_symlink
    # Exercise the refusal on hosts that cannot create native filesystem links.
    monkeypatch.setattr(Path, "is_symlink", lambda path: path == link or original_is_symlink(path))
    outcome = install_skill_dir(linked, source="local", origin="x")
    assert not outcome.imported and "symlink" in outcome.skipped_reason
    assert not (user_dir / "imported").exists()


def test_import_refuses_unconfirmed_dangerous_skills(tmp_path, user_dir) -> None:
    dangerous = _write_skill(tmp_path / "b", "shelly", "Runs shell", extra_frontmatter="required_capabilities: [shell:execute]\n")
    refused = install_skill_dir(dangerous, source="local", origin="x")
    assert not refused.imported and "dangerous" in refused.skipped_reason
    assert install_skill_dir(dangerous, source="local", origin="x", confirm_dangerous=True).imported


def test_import_reports_missing_skills(tmp_path, user_dir) -> None:
    with pytest.raises(SkillImportError):
        import_skills(str(tmp_path / "missing"))
    (tmp_path / "empty").mkdir()
    with pytest.raises(SkillImportError):
        import_skills(str(tmp_path / "empty"))


def test_openjarvis_install_is_an_import_source(tmp_path, user_dir, monkeypatch) -> None:
    monkeypatch.setenv("OPENJARVIS_HOME", str(tmp_path / "oj"))
    _write_skill(tmp_path / "oj" / "skills" / "hermes", "arxiv", "Search arXiv papers")
    _write_skill(tmp_path / "oj" / "skills" / "github", "weather", "Get the weather")

    outcomes = import_skills("openjarvis:weather")

    assert [o.name for o in outcomes] == ["weather"]
    assert outcomes[0].target == user_dir / "imported" / "openjarvis" / "weather"


def test_resolve_source_specs() -> None:
    hermes = resolve_source("hermes:arxiv")
    assert (hermes.kind, hermes.location, hermes.subpath, hermes.name, hermes.label) == ("github", "NousResearch/hermes-agent", "skills", "arxiv", "hermes")
    openclaw = resolve_source("openclaw:alice/todo")
    assert (openclaw.location, openclaw.subpath, openclaw.label) == ("openclaw/skills", "skills/alice/todo", "openclaw")
    url = resolve_source("https://github.com/anthropics/skills/tree/main/skills/docx")
    assert (url.location, url.ref, url.subpath, url.label) == ("anthropics/skills", "main", "skills/docx", "anthropics")
    spec = resolve_source("github:owner/repo/path/to@v1")
    assert (spec.location, spec.subpath, spec.ref) == ("owner/repo", "path/to", "v1")
    with pytest.raises(SkillImportError):
        resolve_source("openclaw:just-a-name")


# Structured runs ------------------------------------------------------------------


def test_extract_json_object_tolerates_fences_and_chatter() -> None:
    assert extract_json_object('```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json_object('Here you go: {"a": {"b": 2}} done') == {"a": {"b": 2}}
    with pytest.raises(ValueError):
        extract_json_object("no json here")


async def test_spec_generation_retries_with_errors_then_escalates(user_dir) -> None:
    entry = SkillRegistry.discover().get("xlsx")
    schema = json.loads((entry.manifest.directory / "schema.json").read_text(encoding="utf-8"))
    calls = []
    replies = ["not json", json.dumps({"sheets": [{"name": "S"}]}), json.dumps(_example_spec("xlsx"))]

    async def _generate(system, user, force):
        calls.append((user, force))
        return replies[len(calls) - 1]

    spec = await generate_valid_spec(entry, schema, "budget", _generate)

    assert spec == _example_spec("xlsx")
    assert [force for _, force in calls] == [False, False, True]
    assert "No JSON object found" in calls[1][0]
    assert "'columns' is a required property" in calls[2][0]


async def test_unapproved_imported_structured_skill_does_not_run(tmp_path, user_dir) -> None:
    source = tmp_path / "src" / "pptx-lite"
    source.mkdir(parents=True)
    (source / "SKILL.md").write_text(
        "---\nname: pptx-lite\ndescription: Lite decks\nmetadata:\n  jarvis:\n    kind: structured\n"
        "    entrypoint: scripts/build.py\n    output_extension: .pptx\n---\nWrite {\"title\": ...}\n",
        encoding="utf-8",
    )
    (source / "schema.json").write_text('{"type": "object"}', encoding="utf-8")
    (source / "scripts").mkdir()
    (source / "scripts" / "build.py").write_text("raise SystemExit(0)\n", encoding="utf-8")
    import_skills(str(source))
    entry = SkillRegistry.discover().get("pptx-lite")

    async def _generate(system, user, force):
        return '{"title": "x"}'

    result = await run_structured_skill(entry, "make a lite deck", generate=_generate)
    assert result["success"] is False and "without approving its scripts" in result["error"]


def test_resolve_output_path(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.setenv("JARVIS_OUTPUT_DIR", str(tmp_path / "jarvis-out"))
    (tmp_path / "Desktop").mkdir()
    project = tmp_path / "project"
    project.mkdir()

    assert resolve_output_path("make slides", "Solar Energy: Basics!", ".pptx") == tmp_path / "jarvis-out" / "Solar-Energy-Basics.pptx"
    assert resolve_output_path("save the deck on my desktop", "Deck", ".pptx") == tmp_path / "Desktop" / "Deck.pptx"
    assert resolve_output_path(f"put it in {project}", "Deck", ".pptx") == project / "Deck.pptx"
    assert resolve_output_path(f"save as {project / 'final.pptx'}", "Deck", ".pptx") == project / "final.pptx"
    (project / "Deck.pptx").write_bytes(b"")
    assert resolve_output_path(f"put it in {project}", "Deck", ".pptx") == project / "Deck-2.pptx"


# CLI agent integration ----------------------------------------------------------------


async def test_cli_agent_runs_matching_document_skill(tmp_path, monkeypatch, user_dir) -> None:
    monkeypatch.setenv("JARVIS_OUTPUT_DIR", str(tmp_path / "out"))

    async def _generate(system, user, force):
        return json.dumps(_example_spec("pptx"))

    statuses = []

    async def _status(text):
        statuses.append(text)

    result = await CLIAgent(spec_generator=_generate).execute("make a presentation on solar energy", status_callback=_status)

    assert result["success"] is True and result["skill"] == "pptx", result
    assert Path(result["output_file_path"]).exists()
    assert statuses == ["Using the pptx skill..."]


def test_cli_prompt_offers_relevant_and_named_skills(user_dir) -> None:
    _write_skill(user_dir, "arxiv", "Search arXiv papers by keyword", body="Call the arXiv API.")
    agent = CLIAgent()

    catalog_prompt = agent._prepare_cli_task("find arxiv papers about transformers")
    assert "Installed skills that may help" in catalog_prompt
    assert "- arxiv: Search arXiv papers by keyword (read " in catalog_prompt
    assert "Call the arXiv API." not in catalog_prompt

    named_prompt = agent._prepare_cli_task("use the arxiv skill to find papers on RAG")
    assert "The user asked you to use the 'arxiv' skill" in named_prompt
    assert "Call the arXiv API." in named_prompt

    assert "skill" not in agent._prepare_cli_task("open notepad").lower()
