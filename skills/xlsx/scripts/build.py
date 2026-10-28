"""
Build an Excel workbook from an xlsx skill spec.

Usage: python build.py --spec spec.json --output budget.xlsx
Prints {"path": ..., "sheets": N} on success. Exit code 2 means the spec is invalid.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
MAX_COLUMN_WIDTH = 60


class SpecError(ValueError):
    pass


def _column_index(columns: list[str], name: str, sheet: str) -> int:
    try:
        return columns.index(name) + 1
    except ValueError:
        raise SpecError(f"Sheet '{sheet}' has no column named {name!r}") from None


def _numeric_columns(rows: list[list], count: int) -> list[int]:
    numeric = []
    for col in range(count):
        values = [row[col] for row in rows if col < len(row) and row[col] not in (None, "")]
        # Formula columns (e.g. "=B2-C2") are summed too; text formulas are rare in totals.
        if values and all(
            (isinstance(v, (int, float)) and not isinstance(v, bool)) or (isinstance(v, str) and v.startswith("="))
            for v in values
        ):
            numeric.append(col + 1)
    return numeric


def _add_chart(ws, sheet: dict, columns: list[str], last_data_row: int) -> None:
    chart_spec = sheet["chart"]
    kind = chart_spec["type"]
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    chart.title = chart_spec.get("title") or None
    category_col = _column_index(columns, chart_spec["category_column"], sheet["name"])
    categories = Reference(ws, min_col=category_col, min_row=2, max_row=last_data_row)
    value_columns = chart_spec["value_columns"][:1] if kind == "pie" else chart_spec["value_columns"]
    for name in value_columns:
        col = _column_index(columns, name, sheet["name"])
        chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=last_data_row), titles_from_data=True)
    chart.set_categories(categories)
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, f"{get_column_letter(len(columns) + 2)}2")


def build(spec: dict, output: Path) -> int:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in spec["sheets"]:
        columns = [str(c) for c in sheet["columns"]]
        rows = sheet.get("rows") or []
        ws = workbook.create_sheet(title=sheet["name"])
        ws.append(columns)
        for col in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([row[i] if i < len(row) else None for i in range(len(columns))])
        last_data_row = len(rows) + 1

        if sheet.get("totals_row") and rows:
            total_row = last_data_row + 1
            ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
            for col in _numeric_columns(rows, len(columns)):
                if col == 1:
                    continue
                letter = get_column_letter(col)
                ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{last_data_row})").font = Font(bold=True)
            for col in range(1, len(columns) + 1):
                ws.cell(row=total_row, column=col).fill = TOTAL_FILL

        for name, fmt in (sheet.get("number_formats") or {}).items():
            col = _column_index(columns, name, sheet["name"])
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row=row_index, column=col).number_format = fmt

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(columns)):
            for cell in row:
                cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

        for col in range(1, len(columns) + 1):
            values = [ws.cell(row=r, column=col).value for r in range(1, ws.max_row + 1)]
            longest = max((len(str(v)) for v in values if v is not None and not str(v).startswith("=")), default=8)
            ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 2, 10), MAX_COLUMN_WIDTH)

        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_data_row}"
        if sheet.get("chart") and rows:
            _add_chart(ws, sheet, columns, last_data_row)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output))
    return len(spec["sheets"])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    output = Path(args.output)
    try:
        count = build(spec, output)
    except (SpecError, KeyError, TypeError, ValueError) as exc:
        print(f"Invalid spec: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"path": str(output.resolve()), "sheets": count}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
